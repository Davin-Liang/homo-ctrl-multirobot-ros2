"""Tests for the Markdown-driven Word report builder."""

import importlib.util
import sys
import unittest
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document


REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = REPOSITORY_ROOT / "docs/reports/2026-09-15-real-robot-tracking"
BUILDER_PATH = REPOSITORY_ROOT / "homo_multirobot_formation_control/scripts/build_real_tracking_deliverables.py"


def load_builder():
    spec = importlib.util.spec_from_file_location("report_builder", BUILDER_PATH)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


class MarkdownWordReportTest(unittest.TestCase):
    def test_markdown_parser_reads_current_sections_images_and_tables(self):
        builder = load_builder()

        blocks = builder.parse_markdown_report(REPORT_DIR / "report_content.md")

        self.assertEqual(blocks[0]["kind"], "title")
        self.assertEqual(blocks[0]["text"], "双全向移动机器人 Leader–Follower 编队跟踪实物实验报告")
        self.assertEqual(
            [block["text"] for block in blocks if block["kind"] == "heading"],
            [
                "1. 实验任务说明", "2. 数据流与控制方法", "3. 实物实验平台",
                "4. 实验过程与统一条件", "5. 实验一：Artstein-HPC 的 Leader 命令速度前馈开/关对比",
                "6. 实验二：Artstein-HPC 与 Artstein-LPC 的阶段性结果",
            ],
        )
        images = [block["path"].name for block in blocks if block["kind"] == "image"]
        images.extend(
            image.name for block in blocks if block["kind"] == "image_table"
            for image in block["images"]
        )
        self.assertIn("omni_robot.jpg", images)
        self.assertIn("experimental_site.jpg", images)
        self.assertIn("hpc_lpc_y_position.png", images)
        self.assertGreaterEqual(len([block for block in blocks if block["kind"] == "table"]), 3)

    def test_word_is_built_from_markdown_and_embeds_current_assets(self):
        builder = load_builder()
        with TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "report.docx"
            builder.build_word(REPORT_DIR, output)

            document = Document(output)
            section = document.sections[0]
            self.assertAlmostEqual(section.page_width.cm, 21.0, places=1)
            self.assertAlmostEqual(section.page_height.cm, 29.7, places=1)
            text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            )
            self.assertIn("双全向移动机器人 Leader–Follower 编队跟踪实物实验报告", text)
            self.assertIn("Artstein 输入时滞补偿", text)
            self.assertIn("末 10 s 距离误差标准差", text)
            self.assertNotIn("视频", text)

            with zipfile.ZipFile(output) as archive:
                media = [name for name in archive.namelist() if name.startswith("word/media/")]
            self.assertGreaterEqual(len(media), 10)

    def test_artstein_original_4d_comparison_can_be_appended_to_a_word_report(self):
        builder = load_builder()
        with TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "report.docx"
            builder.build_word(REPORT_DIR, output)
            builder.append_artstein_original_4d_comparison(REPORT_DIR, output)
            builder.append_artstein_original_4d_comparison(REPORT_DIR, output)
            builder.append_experiment_summary(output)
            builder.append_experiment_summary(output)

            document = Document(output)
            text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            )
            self.assertIn("7. 实验三：Artstein 4D 与原始4D控制器实物对比", text)
            self.assertIn("原始4D控制器（未进行Artstein时滞补偿）", text)
            self.assertIn("0.0722", text)
            self.assertEqual(text.count("7. 实验三：Artstein 4D 与原始4D控制器实物对比"), 1)
            self.assertIn("初始车间距分别为 1.9112 m 和 1.5333 m", text)
            self.assertIn("不宜作为两类控制器性能优劣的主要依据", text)
            self.assertIn("状态预测与时滞补偿对编队稳定性具有积极作用", text)
            self.assertIn("8. 本阶段实验小结", text)
            self.assertEqual(text.count("8. 本阶段实验小结"), 1)


if __name__ == "__main__":
    unittest.main()
