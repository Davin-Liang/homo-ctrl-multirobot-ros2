#!/usr/bin/env python3
"""Build the Chinese Word summary for the 2026-09-15 real-robot study."""

import argparse
import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt as PptPt


LPC_SAFETY_BOUNDARY = (
    "LPC 在 initial_min_lambda=1.5、Leader 速度 0.25 m/s 的实物工况下发生碰撞，"
    "未形成可用于轨迹统计的有效记录。LPC 的 λ初值=2.0 记录同时将 Leader 速度降至"
    " 0.20 m/s；λ初值=2.5 记录使用 0.25 m/s。因此现有 HPC/LPC 结果用于说明阶段性"
    "可运行性和现象，不用于宣称严格单变量条件下的性能优劣。"
)
BANNED_TERMS = ("视频", "二维码", "附件", "链接")
ASSET_NAMES = (
    "control_pipeline.png",
    "feedforward_comparison.png",
    "feedforward_distance_error.png",
    "hpc_lpc_trajectory.png",
    "hpc_lpc_distance_error.png",
)
RADIUS_BASIS = "评价采用控制器 radius=1.0 m。"
COMMON_CONDITIONS = "共同配置：Leader /robot1；Follower /robot2；20 Hz；tau=0.43 s；Td=0.22 s。"
EXPERIMENT_PROCEDURE = (
    "复现实验流程：准备机器人与动捕，核对状态话题和控制参数 → 启动 Follower 控制器与 "
    "Leader 轨迹 → 录制 45 s 并保存 raw.csv、metadata.yaml → 停止轨迹和控制器，确认两车停止。"
)
SOURCE_IDS = {
    "hpc_feedforward_on": "A",
    "hpc_feedforward_off": "B",
    "lpc_lambda_20_v020": "C",
    "lpc_lambda_25_v025": "D",
}


def report_paragraphs():
    """Return body text so factual boundaries can be tested without DOCX parsing."""
    return [
        "双移动机器人 Leader–Follower 实物跟踪实验总结",
        "1. 实验任务说明",
        "本实验在双移动机器人 Leader–Follower 场景中，记录并整理 Artstein-HPC 与 "
        "Artstein-LPC 的阶段性实物跟踪结果。评价半径固定为 1.0 m，所有统计均以该理想"
        "编队半径为唯一评价基准；距离误差定义为实际两车间距减去 1.0 m，表中的末帧误差"
        "为其绝对值。" + RADIUS_BASIS,
        "2. 实物实验平台、数据流与控制方法",
        "两台 mini_omni 全向移动机器人采用动捕状态源，Leader 的轨迹由速度命令驱动，"
        "Follower 依据相对编队误差输出速度命令。控制频率为 20 Hz，记录时长约 45 s。",
        "Follower 的前向预测使用 Follower 当前动捕状态、最近输出命令和电机时间常数 tau "
        "预测其短时状态，以补偿执行器响应。其原理是 Artstein 积分项将历史控制输入纳入状态，"
        "先补偿输入纯滞后，再将状态前向预测到控制时刻，使 HPC/LPC 针对预测状态而非滞后测量状态"
        "计算控制量；它与 Leader 命令速度前馈是两条不同支路。"
        "Leader 命令速度前馈为可选的 Leader cmd_vel 速度增量支路，直接叠加到控制输出，"
        "不是 Follower 前向预测的开关。",
        "3. 实验过程与统一条件",
        "四组有效记录均来自实物平台和动捕状态源。统计只使用条件表与指标表中理想半径为 "
        "1.0 m 的数据，不根据原始目录名或元数据中的控制器字符串重新推断前馈状态。"
        + COMMON_CONDITIONS + EXPERIMENT_PROCEDURE,
        "4. 实验一：Artstein-HPC 的 Leader 命令速度前馈开/关对比",
        "在两组约 45 s、Leader 平均速度约 0.246 m/s 的 Artstein-HPC 实测记录中，"
        "开启 Leader 命令速度前馈的平均绝对距离误差为 0.1202 m、RMS 距离误差为 "
        "0.2791 m；关闭时分别为 0.1282 m 和 0.2895 m。开启组在这两个汇总指标上较低。"
        "末帧绝对距离误差则为 0.0355 m，高于关闭组的 0.0165 m。因此该组对比仅描述"
        "已记录的实测差异，不将其概括为所有指标的一致改善。",
        "5. 实验二：Artstein-HPC 与 Artstein-LPC 的阶段性结果",
        "LPC 两条有效记录的平均绝对距离误差分别为 0.4001 m 与 0.4144 m，RMS 距离"
        "误差分别为 0.4547 m 与 0.4771 m。对应的 initial_min_lambda 与 Leader 速度"
        "同时变化，结果只作为阶段性观察。",
        LPC_SAFETY_BOUNDARY,
        "6. 阶段性成果、当前问题与下一步工作",
        "已形成统一的实物实验条件表、1.0 m 基准指标表和可复核的图表总结。下一步应在"
        "固定 Leader 速度、初始 λ 和其他运行条件后，补充可重复试验；同时继续核对前馈"
        "支路标注与原始实验配置的一致性。",
    ]


def load_report_data(input_dir):
    """Load the committed report inputs and reject a non-1.0 m metric basis."""
    input_dir = Path(input_dir)
    with (input_dir / "experiment-conditions.csv").open(encoding="utf-8-sig", newline="") as stream:
        conditions = list(csv.DictReader(stream))
    with (input_dir / "metrics.csv").open(encoding="utf-8-sig", newline="") as stream:
        metrics = list(csv.DictReader(stream))
    if not conditions or not metrics:
        raise ValueError("条件表和指标表均不能为空")
    if {row["ideal_radius_m"] for row in metrics} != {"1.0"}:
        raise ValueError("报告只接受理想编队半径为 1.0 m 的指标")
    for asset_name in ASSET_NAMES:
        if not (input_dir / "assets" / asset_name).is_file():
            raise FileNotFoundError(input_dir / "assets" / asset_name)
    return conditions, metrics


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_run_font(run, size=10.0, bold=False):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def add_body_paragraph(document, text, *, bold=False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)
    return paragraph


def add_heading(document, text, level):
    heading = document.add_heading()
    heading.style = document.styles[f"Heading {level}"]
    run = heading.add_run(text)
    set_run_font(run, size=15 if level == 1 else 12, bold=True)
    return heading


def add_caption(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, size=9)
    return paragraph


def add_image(document, image_path, caption):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Cm(22.5))
    add_caption(document, caption)


def add_table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, "D9EAF7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_run_font(run, size=8, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(str(value))
            set_run_font(run, size=7.5)
    document.add_paragraph()
    return table


def metric_rows(metrics, comparison):
    return [
        (
            row["display_label"],
            row["sample_count"],
            f"{float(row['mean_abs_distance_error_m']):.4f}",
            f"{float(row['rms_distance_error_m']):.4f}",
            f"{float(row['final_distance_error_m']):.4f}",
        )
        for row in metrics
        if row["comparison"] == comparison or (
            comparison == "hpc_lpc_stage_result" and row["experiment_id"] == "hpc_feedforward_on")
    ]


IMAGE_PATTERN = re.compile(r"!\[[^]]*\]\(([^)]+)\)")


def _table_cells(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def parse_markdown_report(source_path):
    """Parse the limited Markdown subset used by the editable report source."""
    lines = Path(source_path).read_text(encoding="utf-8").splitlines()
    blocks, index = [], 0
    while index < len(lines):
        line = lines[index].strip()
        if not line or line.startswith("<!--"):
            index += 1
            continue
        if line.startswith("# "):
            blocks.append({"kind": "title", "text": line[2:].strip()})
            index += 1
            continue
        if line.startswith("## "):
            blocks.append({"kind": "heading", "text": line[3:].strip()})
            index += 1
            continue
        image_match = IMAGE_PATTERN.fullmatch(line)
        if image_match:
            blocks.append({"kind": "image", "path": Path(image_match.group(1))})
            index += 1
            continue
        if line.startswith("*") and line.endswith("*"):
            blocks.append({"kind": "caption", "text": line.strip("*").strip()})
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            headers = _table_cells(table_lines[0])
            rows = [_table_cells(table_line) for table_line in table_lines[2:]]
            image_rows = [row for row in rows if any(IMAGE_PATTERN.fullmatch(cell) for cell in row)]
            if image_rows:
                images = [Path(IMAGE_PATTERN.fullmatch(cell).group(1)) for cell in image_rows[0]]
                captions = next((row for row in rows if row is not image_rows[0]), [""] * len(images))
                blocks.append({"kind": "image_table", "headers": headers,
                               "images": images, "captions": captions})
            else:
                blocks.append({"kind": "table", "headers": headers, "rows": rows})
            continue
        paragraph_lines = []
        while index < len(lines):
            candidate = lines[index].strip()
            if (not candidate or candidate.startswith("#") or candidate.startswith("|")
                    or candidate.startswith("<!--") or IMAGE_PATTERN.fullmatch(candidate)
                    or (candidate.startswith("*") and candidate.endswith("*"))):
                break
            paragraph_lines.append(candidate)
            index += 1
        blocks.append({"kind": "paragraph", "text": " ".join(paragraph_lines).replace("`", "")})
    return blocks


def add_image_table(document, headers, images, captions, base_dir):
    table = document.add_table(rows=2, cols=len(images))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(table.rows[0].cells, headers):
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = cell.paragraphs[0].add_run(header)
        set_run_font(run, size=9, bold=True)
    for cell, image, caption in zip(table.rows[1].cells, images, captions):
        image_path = base_dir / image
        if not image_path.is_file():
            raise FileNotFoundError(image_path)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(str(image_path), width=Cm(9.4))
        caption_paragraph = cell.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption_paragraph.add_run(caption.strip("*").strip())
        set_run_font(run, size=8)
    document.add_paragraph()


def build_word(input_dir, word_out):
    """Build the Word report directly from the editable Markdown source."""
    input_dir = Path(input_dir)
    word_out = Path(word_out)
    blocks = parse_markdown_report(input_dir / "report_content.md")
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10)

    for block in blocks:
        kind = block["kind"]
        if kind == "title":
            title = document.add_paragraph()
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title.paragraph_format.space_after = Pt(14)
            title_run = title.add_run(block["text"])
            set_run_font(title_run, size=18, bold=True)
        elif kind == "heading":
            add_heading(document, block["text"], 1)
        elif kind == "paragraph":
            add_body_paragraph(document, block["text"])
        elif kind == "caption":
            add_caption(document, block["text"])
        elif kind == "image":
            image_path = input_dir / block["path"]
            if not image_path.is_file():
                raise FileNotFoundError(image_path)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            width = Cm(10.2) if image_path.suffix.lower() in {".jpg", ".jpeg"} else Cm(22.5)
            paragraph.add_run().add_picture(str(image_path), width=width)
        elif kind == "image_table":
            add_image_table(document, block["headers"], block["images"], block["captions"], input_dir)
        elif kind == "table":
            add_table(document, block["headers"], block["rows"])
    word_out.parent.mkdir(parents=True, exist_ok=True)
    document.save(word_out)
    return word_out


def _load_metric_rows(input_dir):
    with (Path(input_dir) / "metrics.csv").open(encoding="utf-8-sig", newline="") as stream:
        return {row["experiment_id"]: row for row in csv.DictReader(stream)}


def append_artstein_original_4d_comparison(input_dir, word_path):
    """Append the requested Artstein 4D vs original 4D experiment to an existing Word report."""
    input_dir = Path(input_dir)
    word_path = Path(word_path)
    metrics = _load_metric_rows(input_dir)
    required = ("hpc_lpc_reference", "original_4d_reference")
    missing = [experiment_id for experiment_id in required if experiment_id not in metrics]
    if missing:
        raise ValueError("缺少实验三指标：" + ", ".join(missing))
    document = Document(word_path)
    heading_text = "7. 实验三：Artstein 4D 与原始4D控制器实物对比"
    for paragraph in document.paragraphs:
        if paragraph.text == heading_text:
            body = paragraph._p.getparent()
            start = list(body).index(paragraph._p)
            for element in list(body)[start:]:
                if element.tag.endswith("sectPr"):
                    break
                body.remove(element)
            break
    add_heading(document, heading_text, 1)
    add_body_paragraph(
        document,
        "本实验比较采用 Artstein 输入时滞补偿与 Follower 状态前向预测的 4D 控制器，"
        "以及未使用 Artstein 时滞补偿的原始4D控制器。两组记录均来自动捕实物实验，"
        "Leader 平均速度分别约为 0.252 m/s 和 0.250 m/s。",
    )
    images = (
        ("artstein_original_4d_trajectory.png", "图 11　Artstein 4D 与原始4D控制器轨迹对比"),
        ("artstein_original_4d_velocity_components.png", "图 12　Artstein 4D 与原始4D控制器 Vx/Vy 速度对比"),
        ("artstein_original_4d_x_position.png", "图 13　Artstein 4D 与原始4D控制器 X 坐标对比"),
        ("artstein_original_4d_y_position.png", "图 14　Artstein 4D 与原始4D控制器 Y 坐标对比"),
    )
    for filename, caption in images:
        path = input_dir / "assets" / filename
        if not path.is_file():
            raise FileNotFoundError(path)
        add_image(document, path, caption)
    labels = {
        "hpc_lpc_reference": "Artstein 4D（Artstein时滞补偿）",
        "original_4d_reference": "原始4D控制器（未进行Artstein时滞补偿）",
    }
    rows = []
    for experiment_id in required:
        row = metrics[experiment_id]
        rows.append((
            labels[experiment_id],
            f"{float(row['mean_abs_distance_error_m']):.4f}",
            f"{float(row['rms_distance_error_m']):.4f}",
            f"{float(row['tail_mean_abs_distance_error_m']):.4f}",
            f"{float(row['tail_distance_error_std_m']):.4f}",
            f"{float(row['mean_relative_velocity_error_mps']):.4f}",
            f"{float(row['tail_mean_relative_velocity_error_mps']):.4f}",
        ))
    add_table(document, (
        "工况", "平均距离误差 (m)", "RMS 距离误差 (m)", "末 10 s 距离误差 (m)",
        "末 10 s 距离误差标准差 (m)", "平均相对速度误差 (m/s)",
        "末 10 s 相对速度误差 (m/s)"), rows)
    add_caption(document, "表 4　Artstein 4D 与原始4D控制器的实物跟踪指标对比")
    add_body_paragraph(
        document,
        "由表 4 可见，在本次记录中，原始4D控制器的平均距离误差和 RMS 距离误差"
        "小于 Artstein 4D 组；Artstein 4D 组的末 10 s 距离误差标准差则略小。"
        "该对比反映两组实物记录下的实际结果，后续需要在统一初始状态和重复试验条件下"
        "进一步分析延迟补偿对不同工况的影响。",
    )
    document.save(word_path)
    return word_path


def _add_ppt_text(slide, text, left, top, width, height, *, size=20, bold=False,
                  color=(31, 78, 121), align=PP_ALIGN.LEFT, wrap=False):
    """Add an editable text box to a presentation slide."""
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.word_wrap = wrap
    paragraph = box.text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "宋体"
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return box


def _add_ppt_title(slide, title):
    _add_ppt_text(slide, title, 0.55, 0.28, 12.1, 0.48, size=26, bold=True)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.87),
                                    Inches(1.55), Inches(0.05))
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor(0, 112, 192)
    accent.line.fill.background()


def _add_ppt_bullets(slide, lines, *, top=1.2, font_size=18):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.7), Inches(5.7))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.level = 0
        paragraph.font.name = "宋体"
        paragraph.font.size = PptPt(font_size)
        paragraph.font.color.rgb = RGBColor(45, 45, 45)
        paragraph.space_after = PptPt(14)
    return box


def _add_ppt_image(slide, image_path, left, top, width, height):
    slide.shapes.add_picture(str(image_path), Inches(left), Inches(top),
                             width=Inches(width), height=Inches(height))


def build_presentation(input_dir, ppt_out):
    """Build the nine-slide editable PowerPoint summary without media content."""
    input_dir = Path(input_dir)
    ppt_out = Path(ppt_out)
    conditions, metrics = load_report_data(input_dir)
    metric_by_id = {row["experiment_id"]: row for row in metrics}
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    # 1. Title
    slide = presentation.slides.add_slide(blank)
    _add_ppt_text(slide, "双移动机器人实物跟踪实验汇报", 0.8, 2.1, 11.8, 0.7,
                  size=32, bold=True, align=PP_ALIGN.CENTER)
    _add_ppt_text(slide, "Artstein-HPC / Artstein-LPC 阶段性实物结果", 1.1, 3.0, 11.1, 0.45,
                  size=20, color=(80, 80, 80), align=PP_ALIGN.CENTER)
    _add_ppt_text(slide, "统一评价基准：理想编队半径 1.0 m", 1.1, 4.05, 11.1, 0.4,
                  size=18, bold=True, color=(0, 112, 192), align=PP_ALIGN.CENTER)

    # 2. Task and objectives
    slide = presentation.slides.add_slide(blank)
    _add_ppt_title(slide, "实验任务与阶段目标")
    _add_ppt_bullets(slide, [
        "在双 mini_omni Leader–Follower 场景中记录 Artstein-HPC 与 Artstein-LPC 实物跟踪。",
        "统一评价：实际两车间距 − 1.0 m；末帧指标取该距离误差的绝对值。",
        "本阶段目标：形成可复核的条件、指标与图表，并明确现有结果的适用边界。",
    ])

    # 3. Platform and data flow
    slide = presentation.slides.add_slide(blank)
    _add_ppt_title(slide, "实物平台与数据流")
    _add_ppt_image(slide, input_dir / "assets" / "control_pipeline.png", 0.8, 1.15, 7.3, 5.5)
    _add_ppt_bullets(slide, [
        "两台 mini_omni 全向移动机器人；状态源为动捕。",
        "Leader 由速度命令驱动；Follower 依据相对编队误差输出速度命令。",
        "控制频率 20 Hz；单条有效记录约 45 s。",
    ], top=1.45, font_size=16).left = Inches(8.35)
    # Reset the bullet box geometry after reuse of the compact helper.
    slide.shapes[-1].width = Inches(4.3)
    slide.shapes[-1].height = Inches(4.8)

    # 4. Method
    slide = presentation.slides.add_slide(blank)
    _add_ppt_title(slide, "控制方法：Artstein + Follower 前向预测 + HPC/LPC")
    _add_ppt_bullets(slide, [
        "Artstein 积分项纳入历史控制输入以补偿纯滞后；HPC/LPC 为两类控制配置。",
        "Follower 前向预测：由 Follower 当前动捕状态、最近输出命令和电机时间常数 τ 预测短时状态，补偿执行器响应。",
        "控制律针对预测状态而非滞后测量状态计算控制量。",
        "Leader 命令速度前馈：可选 Leader cmd_vel 速度增量支路，直接叠加到控制输出。",
        "两者是不同支路：前馈开/关不是 Follower 前向预测开/关。",
    ], font_size=17)

    # 5. Conditions
    slide = presentation.slides.add_slide(blank)
    _add_ppt_title(slide, "实验流程与统一条件")
    _add_ppt_bullets(slide, [
        COMMON_CONDITIONS + "状态源：动捕。",
        RADIUS_BASIS,
        EXPERIMENT_PROCEDURE,
        "HPC 前馈开/关：Leader 均速约 0.246 m/s。LPC：λ=2.0 / 0.20 m/s；λ=2.5 / 0.25 m/s。",
        "前馈开关按元数据 enable_leader_cmd_feedforward 核对，不依据目录名或 controller 字符串。",
        "后续图表来源目录相对于 homo_multirobot_formation_control；使用各目录 raw.csv 与 metadata.yaml。",
    ], font_size=15)

    # 6. Feedforward experiment
    on = metric_by_id["hpc_feedforward_on"]
    off = metric_by_id["hpc_feedforward_off"]
    slide = presentation.slides.add_slide(blank)
    _add_ppt_title(slide, "实验一：Leader 命令速度前馈开/关")
    _add_ppt_image(slide, input_dir / "assets" / "feedforward_comparison.png", 0.55, 1.12, 5.9, 4.15)
    _add_ppt_image(slide, input_dir / "assets" / "feedforward_distance_error.png", 6.75, 1.12, 5.9, 4.15)
    _add_ppt_text(slide,
                  f"开启：平均绝对误差 {float(on['mean_abs_distance_error_m']):.4f} m；RMS {float(on['rms_distance_error_m']):.4f} m；末帧 {float(on['final_distance_error_m']):.4f} m\n"
                  f"关闭：平均绝对误差 {float(off['mean_abs_distance_error_m']):.4f} m；RMS {float(off['rms_distance_error_m']):.4f} m；末帧 {float(off['final_distance_error_m']):.4f} m\n"
                  "开启组前两项汇总指标较低，但末帧绝对误差更高；仅描述已记录实测差异。",
                  0.72, 5.55, 11.9, 1.0, size=15, color=(45, 45, 45), wrap=True)
    _add_ppt_text(slide, "两图来源 A、B（A：HPC 前馈开；B：HPC 前馈关）\n" + "\n".join(
        f"{SOURCE_IDS[row['experiment_id']]}：{row['source_dir']} / {row['trial_id']}"
        for row in (on, off)), 0.72, 6.65, 11.9, 0.65, size=10, color=(80, 80, 80))

    # 7. HPC/LPC result and caveat
    lpc20 = metric_by_id["lpc_lambda_20_v020"]
    lpc25 = metric_by_id["lpc_lambda_25_v025"]
    slide = presentation.slides.add_slide(blank)
    _add_ppt_title(slide, "实验二：HPC/LPC 阶段性结果与安全边界")
    _add_ppt_image(slide, input_dir / "assets" / "hpc_lpc_trajectory.png", 0.55, 1.1, 5.9, 3.5)
    _add_ppt_image(slide, input_dir / "assets" / "hpc_lpc_distance_error.png", 6.75, 1.1, 5.9, 3.5)
    _add_ppt_text(slide,
                  f"HPC 代表组 A（前馈开）：平均绝对误差 {float(on['mean_abs_distance_error_m']):.4f} m，RMS {float(on['rms_distance_error_m']):.4f} m。\n"
                  f"LPC λ=2.0：平均绝对误差 {float(lpc20['mean_abs_distance_error_m']):.4f} m，RMS {float(lpc20['rms_distance_error_m']):.4f} m；"
                  f"LPC λ=2.5：{float(lpc25['mean_abs_distance_error_m']):.4f} m，{float(lpc25['rms_distance_error_m']):.4f} m。\n"
                  "安全边界：λ初值=1.5、Leader 速度 0.25 m/s 时发生碰撞，未形成有效轨迹统计。λ初值与 Leader 速度同时变化，"
                  "现有 HPC/LPC 结果仅说明阶段性可运行性和现象，不宣称严格单变量性能优劣。",
                  0.72, 4.78, 11.9, 1.5, size=14, color=(45, 45, 45), wrap=True)
    _add_ppt_text(slide, "两图来源 A、C、D（A：HPC 前馈开；C：LPC λ=2.0；D：LPC λ=2.5）\n" + "\n".join(
        f"{SOURCE_IDS[row['experiment_id']]}：{row['source_dir']} / {row['trial_id']}"
        for row in (on, lpc20, lpc25)), 0.72, 6.4, 11.9, 0.85, size=10, color=(80, 80, 80))

    # 8. Manual video placeholder only: shapes and text, no media or relationships.
    slide = presentation.slides.add_slide(blank)
    _add_ppt_title(slide, "最新实物实验视频（占位页）")
    placeholder = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.866667), Inches(1.20),
                                         Inches(9.6), Inches(5.4))
    placeholder.fill.background()
    placeholder.line.color.rgb = RGBColor(0, 112, 192)
    placeholder.line.width = PptPt(2.25)
    _add_ppt_text(slide, "请在此处手动插入最新实物实验视频", 1.35, 3.1, 10.6, 0.45,
                  size=24, bold=True, align=PP_ALIGN.CENTER)
    _add_ppt_text(slide, "建议对应：Artstein-LPC，λ初值=2.5，Leader 速度=0.25 m/s",
                  1.35, 3.75, 10.6, 0.35, size=16, color=(80, 80, 80), align=PP_ALIGN.CENTER)

    # 9. Conclusion
    slide = presentation.slides.add_slide(blank)
    _add_ppt_title(slide, "阶段结论与下一步工作")
    _add_ppt_bullets(slide, [
        "已形成统一的实物条件表、1.0 m 基准指标表及可复核图表。",
        "HPC 前馈开/关结果存在指标间差异，应保留原始实测描述，不外推为全面改善。",
        "LPC 结果受 λ初值与 Leader 速度共同变化以及 λ=1.5 碰撞边界限制，仅作阶段性观察。",
        "下一步：固定 Leader 速度、初始 λ 及其他运行条件，补充可重复试验，并持续核对前馈支路标注。",
    ], font_size=17)

    ppt_out.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(ppt_out)
    return ppt_out


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--word-out", type=Path)
    parser.add_argument("--ppt-out", type=Path)
    parser.add_argument("--skip-word", action="store_true")
    parser.add_argument("--skip-ppt", action="store_true")
    return parser.parse_args()


def main():
    args = parse_args()
    if args.skip_word and args.skip_ppt:
        raise ValueError("不能同时跳过 Word 和 PPT")
    if not args.skip_word:
        if args.word_out is None:
            raise ValueError("生成 Word 时必须提供 --word-out")
        print(f"已生成 Word：{build_word(args.input, args.word_out)}")
    if not args.skip_ppt:
        if args.ppt_out is None:
            raise ValueError("生成 PPT 时必须提供 --ppt-out")
        print(f"已生成 PPT：{build_presentation(args.input, args.ppt_out)}")


if __name__ == "__main__":
    main()
